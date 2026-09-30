"""Form ADV landings for stage ``thirteenf_filer_type`` (lane OWN, S5.1). See docs/ALPHA_PANEL_OWNERSHIP.md.

Sources (SEC hosts only, approved agent and host-wide 5 req/s limiter via :func:`atx_db.sec_http.sec_session`):

* **Monthly firm rosters** (``ROSTER_PAGE``, "Information About Registered Investment Advisers and Exempt Reporting
  Advisers"): one file per month and kind (``registered`` = SEC-registered advisers, ``exempt`` = exempt reporting
  advisers), a zip holding one CSV or XLSX, or a bare XLSX. Each is a snapshot of every adviser's current Form ADV
  Part 1A answers on the file date (CRD, SEC 801-/802- number, names, main office, Item 5.D client types with
  regulatory AUM, 5.F RAUM, 5.G services, 6.A other businesses, 7.B private funds; from 2023 also the adviser's
  EDGAR CIK and private-fund counts by type). ``fetch-rosters`` lands one file at a time into
  ``data/raw/sec_adv/rosters/``, parses the kept columns (:data:`ROSTER_COLUMNS`) into
  ``thirteenf_filer_type/adv/rosters/<date>_<kind>.parquet`` and deletes the file; ``receipts.jsonl`` keeps url,
  bytes, sha256, http status, Last-Modified and the header digest.
* **Form ADV Part 1 bulk tables** 2011-11-05 .. 2024-12-31 (``ADV_DATA_PAGE``, two zips of 0.70 / 0.43 GB). Only
  the members in :data:`BULK_MEMBERS` are read: one ranged GET of the zip tail gives the central directory, one
  ranged GET per member returns its local header and deflated bytes (the rest of the zip is never downloaded). The
  range is landed as served into ``data/raw/sec_adv/bulk/``, inflated with a CRC-32 check, parsed into
  ``thirteenf_filer_type/adv/bulk/<table>.parquet`` (all columns VARCHAR, names snake-cased) and deleted. The SEC
  posts no bulk Part 1 data after 2024-12-31 (the page points to IAPD, which has no bulk download).

Clocks:
* ``adv-roster-publication-v1``: a roster dated D is available at its HTTP Last-Modified when that lies within
  [D, D + 10 days] (recent files are posted the same day: ia07012026 Last-Modified 2026-07-01 11:41 UTC); otherwise
  the file was re-posted later and ``available_at`` = D + 14 days 00:00 UTC (documented floor, ``vintage_risk``).
* ``adv-bulk-publication-v1``: bulk rows are available at the zip's Last-Modified (part 1 2026-05-01, part 2
  2025-09-26). The ADV filings behind them were public on IAPD when filed, but not in bulk; every bulk row carries
  ``vintage_risk = 'adv_bulk_backfill'`` and its own filing date as ``evidence_date``.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import io
import json
import os
import re
import shutil
import sys
import time
import zipfile
from pathlib import Path
from typing import Any
from urllib.parse import urljoin

from . import common as C
from . import sec_docs as D
from . import shortflow_common as S

STAGE = "thirteenf_filer_type"
RAW_DIR = S.RAW_ROOT / "sec_adv"
ROSTER_PAGE = ("https://www.sec.gov/data-research/sec-markets-data/"
               "information-about-registered-investment-advisers-exempt-reporting-advisers")
ADV_DATA_PAGE = "https://www.sec.gov/foia-services/frequently-requested-documents/form-adv-data"
BULK_ZIPS = {
    "part1": "https://www.sec.gov/files/adv-filing-data-20111105-20241231-part1.zip",
    "part2": "https://www.sec.gov/files/adv-filing-data-20111105-20241231-part2.zip",
}
#: bulk table -> (zip part, member base name without the date suffix)
BULK_MEMBERS = {
    "adv_filing_types": ("part1", "ADV_Filing_Types"),
    "ia_1d3_cik": ("part1", "IA_1D3_CIK"),
    "era_1d3_cik": ("part1", "ERA_1D3_CIK"),
    "ia_schedule_d_7b1": ("part2", "IA_Schedule_D_7B1"),
    "era_schedule_d_7b1": ("part1", "ERA_Schedule_D_7B1"),
    "ia_adv_base_a": ("part1", "IA_ADV_Base_A"),
    "era_adv_base": ("part1", "ERA_ADV_Base"),
}
ROSTER_RULE = "adv-roster-publication-v1"
BULK_RULE = "adv-bulk-publication-v1"
PLAUSIBLE_DAYS = 10
FLOOR_DAYS = 14
MAX_FILE_BYTES = 400 * 1024 ** 2
ROSTER_START = dt.date(2018, 10, 1)
_CLIENT_TYPES = "abcdefghijklmn"
#: roster header -> kept column (headers are matched after collapsing whitespace)
ROSTER_COLUMNS: dict[str, str] = {
    "Organization CRD#": "crd", "SEC#": "sec_number", "Firm Type": "firm_type", "CIK#": "cik",
    "Total number of CIK numbers": "n_cik", "Primary Business Name": "business_name", "Legal Name": "legal_name",
    "Main Office City": "city", "Main Office State": "state", "Main Office Country": "country",
    "Main Office Postal Code": "postal_code", "SEC Current Status": "sec_status",
    "SEC Status Effective Date": "sec_status_date", "Latest ADV Filing Date": "latest_adv_filing_date",
    "Umbrella Registration": "umbrella", "3A": "legal_form",
    "5E(1)": "fee_pct_aum", "5E(6)": "fee_performance", "5B(1)": "n_employees_advisory", "5C(1)": "n_clients",
    "5F(2)(a)": "raum_discretionary", "5F(2)(b)": "raum_nondiscretionary", "5F(2)(c)": "raum_total",
    "5G(2)": "svc_individuals", "5G(3)": "svc_ric", "5G(4)": "svc_piv", "5G(5)": "svc_institutional",
    "6A(1)": "act_broker_dealer", "6A(2)": "act_registered_rep", "6A(3)": "act_cpo_cta", "6A(6)": "act_insurance_agent",
    "6A(7)": "act_bank", "6A(8)": "act_trust_company",
    "7B": "private_fund_adviser", "Count of Private Funds - 7B(1)": "n_private_funds",
    "Any Hedge Funds": "any_hedge_funds", "Total number of Hedge funds": "n_hedge_funds",
    "Total number of PE funds": "n_pe_funds", "Total number of VC funds": "n_vc_funds",
    "Total number of Real Estate funds": "n_real_estate_funds", "Total number of Securitized funds": "n_securitized_funds",
    "Total number of Liquidity funds": "n_liquidity_funds", "Total number of Other funds": "n_other_funds",
    "Total Gross Assets of Private Funds": "private_fund_gav",
    **{f"5D({x})(1)": f"clients_{x}" for x in _CLIENT_TYPES},
    **{f"5D({x})(3)": f"raum_{x}" for x in _CLIENT_TYPES},
}
ROSTER_FIELDS = tuple(dict.fromkeys(ROSTER_COLUMNS.values()))
_NAME_RE = re.compile(r"^(?:ia)?(\d{8}|\d{6})", re.IGNORECASE)


# ---------------------------------------------------------------- pure helpers
def roster_date(name: str) -> dt.date | None:
    """File name -> roster date: ``ia07012026.zip`` (MMDDYYYY), ``ia070122.zip`` / ``010118-exempt.zip`` (MMDDYY)."""
    base = name.rsplit("/", 1)[-1]
    m = _NAME_RE.match(base)
    if not m:
        return None
    s = m.group(1)
    try:
        if len(s) == 8:
            return dt.date(int(s[4:]), int(s[:2]), int(s[2:4]))
        return dt.date(2000 + int(s[4:]), int(s[:2]), int(s[2:4]))
    except ValueError:
        return None


def roster_links(html: str, start: dt.date = ROSTER_START, end: dt.date | None = None) -> list[tuple[dt.date, str, str]]:
    """``[(date, kind, url)]`` for every roster zip / xlsx linked on ``ROSTER_PAGE`` in [start, end], oldest first.

    ``kind`` is ``exempt`` when the file name says so, else ``registered``. Several files for one (date, kind)
    (re-posted corrections such as ``ia090117_2.zip``) keep the lexicographically last URL.
    """
    out: dict[tuple[dt.date, str], str] = {}
    for href in re.findall(r"href=[\"']([^\"']+\.(?:zip|xlsx))[\"']", html, re.IGNORECASE):
        base = href.rsplit("/", 1)[-1].lower()
        if "no-data" in base or not (base.startswith("ia") or base[:1].isdigit()):
            continue
        d = roster_date(base)
        if d is None or d < start or (end is not None and d > end):
            continue
        kind = "exempt" if "exempt" in base else "registered"
        url = urljoin(ROSTER_PAGE, href)
        key = (d, kind)
        if key not in out or url > out[key]:
            out[key] = url
    return sorted(((d, k, u) for (d, k), u in out.items()), key=lambda r: (r[0], r[1]))


def roster_available_at(d: dt.date, last_modified: dt.datetime | None) -> tuple[dt.datetime, str, bool]:
    """``(available_at, basis, vintage_risk)`` per ``adv-roster-publication-v1``."""
    if last_modified is not None and S.at_utc(d) <= last_modified <= S.at_utc(d + dt.timedelta(days=PLAUSIBLE_DAYS)):
        return last_modified, "last_modified", False
    return S.at_utc(d + dt.timedelta(days=FLOOR_DAYS)), "date_plus_14d_floor", True


def _norm_header(h: Any) -> str:
    return re.sub(r"\s+", " ", str(h or "")).strip()


def _cell(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()[:10]
    s = str(v).strip()
    return s or None


def roster_rows(header: list[Any], rows: Any) -> tuple[list[dict[str, str | None]], dict[str, Any]]:
    """Keep :data:`ROSTER_COLUMNS` from one roster sheet; returns rows and a header report."""
    hn = [_norm_header(h) for h in header]
    idx = {}
    for i, h in enumerate(hn):
        out = ROSTER_COLUMNS.get(h)
        if out is not None and out not in idx:
            idx[out] = i
    kept = []
    for r in rows:
        if r is None:
            continue
        rec = {k: (_cell(r[i]) if i < len(r) else None) for k, i in idx.items()}
        if rec.get("crd") is None:
            continue
        kept.append(rec)
    report = {"columns": len(hn), "header_sha256": hashlib.sha256("|".join(hn).encode()).hexdigest(),
              "kept": sorted(idx), "missing": sorted(set(ROSTER_FIELDS) - set(idx))}
    return kept, report


def read_roster_file(path: Path) -> tuple[list[dict[str, str | None]], dict[str, Any]]:
    """Roster file on disk (zip with one CSV/XLSX, or XLSX) -> kept rows, streamed (a CSV member is read line by
    line; an XLSX member is extracted next to the file for openpyxl's read-only reader). A zip holding only a notice
    (``.txt``, e.g. the January 2019 shutdown file) -> ``([], {'notice': text})``."""
    name = path.name.lower()
    if name.endswith(".zip"):
        with zipfile.ZipFile(path) as z:
            members = [i for i in z.infolist() if not i.is_dir()]
            data = [i for i in members if i.filename.lower().endswith((".csv", ".xlsx", ".xls"))]
            if not data:
                txt = b" ".join(z.read(i) for i in members)[:500].decode("latin-1", "replace")
                return [], {"notice": txt, "members": [i.filename for i in members]}
            m = max(data, key=lambda i: i.file_size)
            if m.filename.lower().endswith(".csv"):
                with z.open(m) as fh:
                    rd = csv.reader(io.TextIOWrapper(fh, encoding="cp1252", errors="replace", newline=""))
                    header = next(rd)
                    return roster_rows(header, rd)
            inner = path.with_name(path.name + "." + m.filename.rsplit("/", 1)[-1].rsplit(".", 1)[-1])
            with z.open(m) as src, inner.open("wb") as dst:
                shutil.copyfileobj(src, dst, 1 << 20)
            try:
                return read_roster_file(inner)
            finally:
                inner.unlink(missing_ok=True)
    if name.endswith(".csv"):
        with path.open(encoding="cp1252", errors="replace", newline="") as fh:
            rd = csv.reader(fh)
            header = next(rd)
            return roster_rows(header, rd)
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        header = list(next(it))
        return roster_rows(header, it)
    finally:
        wb.close()


def snake(name: str) -> str:
    s = re.sub(r"[^0-9a-zA-Z]+", "_", name.strip()).strip("_").lower()
    return s or "col"


# ---------------------------------------------------------------- landing
def ledger() -> S.Ledger:
    return S.Ledger(RAW_DIR / "receipts.jsonl")


def adv_dir() -> Path:
    p = C.stage_dir(STAGE) / "adv"
    p.mkdir(parents=True, exist_ok=True)
    return p


def _session():
    from atx_db.sec_http import sec_session

    if not _SESSION:
        _SESSION.append(sec_session())
    return _SESSION[0]


_SESSION: list[Any] = []


def _get(url: str, headers: dict[str, str] | None = None, timeout: float = 600.0) -> tuple[int, bytes, dict[str, str]]:
    r = _session().get(url, headers=headers or {}, timeout=timeout)
    try:
        return int(r.status_code), r.content, {k.lower(): v for k, v in r.headers.items()}
    finally:
        r.close()


def fetch_rosters(start: dt.date = ROSTER_START, end: dt.date | None = None, limit: int | None = None,
                  months: tuple[int, ...] | None = None) -> dict[str, Any]:
    """Land, parse and delete every roster in [start, end] not yet parsed (resumable on the ledger).

    ``months`` keeps only rosters dated in those calendar months (the 13F build needs the roster in force at each
    13F deadline, i.e. the February / May / August / November rosters)."""
    led = ledger()
    st, body, _ = _get(ROSTER_PAGE, timeout=120)
    if st != 200:
        raise RuntimeError(f"roster page HTTP {st}")
    led.append({"key": "roster_page", "kind": "page", "url": ROSTER_PAGE, "http_status": st, "bytes": len(body),
                "sha256": hashlib.sha256(body).hexdigest(), "fetched_at": S.utc_now()})
    links = roster_links(body.decode("utf-8", "replace"), start, end)
    if months:
        links = [x for x in links if x[0].month in months]
    have = led.latest()
    out_dir = adv_dir() / "rosters"
    out_dir.mkdir(parents=True, exist_ok=True)
    stats = {"listed": len(links), "parsed_now": 0, "notices": 0, "skipped": 0}
    n = 0
    for d, kind, url in links:
        key = f"roster:{d.isoformat()}:{kind}"
        dest = out_dir / f"{d.isoformat()}_{kind}.parquet"
        prev = have.get(key)
        if prev and prev.get("status") in ("parsed", "notice") and prev.get("url") == url and \
                (prev.get("status") == "notice" or dest.exists()):
            stats["skipped"] += 1
            continue
        if limit is not None and n >= limit:
            break
        n += 1
        S.require_free(1.0)
        t0 = time.perf_counter()
        st, body, hdr = _get(url)
        rec: dict[str, Any] = {"key": key, "kind": "roster", "roster_kind": kind, "roster_date": d.isoformat(), "url": url,
                               "http_status": st, "bytes": len(body), "fetched_at": S.utc_now(),
                               "last_modified": hdr.get("last-modified")}
        if st != 200 or len(body) > MAX_FILE_BYTES:
            rec["status"] = "http_error" if st != 200 else "too_large"
            led.append(rec)
            continue
        rec["sha256"] = hashlib.sha256(body).hexdigest()
        raw = RAW_DIR / "rosters" / url.rsplit("/", 1)[-1]
        raw.parent.mkdir(parents=True, exist_ok=True)
        raw.write_bytes(body)   # as served, deleted once parsed
        del body
        rows, report = read_roster_file(raw)
        if not rows:
            rec.update(status="notice", notice=report.get("notice"), members=report.get("members"))
            raw.unlink()
            led.append(rec)
            stats["notices"] += 1
            continue
        av, basis, vr = roster_available_at(d, S.http_date(hdr.get("last-modified")))
        _write_roster(rows, d, kind, av, basis, vr, url.rsplit("/", 1)[-1], dest)
        raw.unlink()
        rec.update(status="parsed", rows=len(rows), header=report, available_at=av.isoformat(), available_basis=basis,
                   vintage_risk=vr, parquet=dest.name, parquet_sha256=C.sha256_file(dest), raw_deleted=True,
                   elapsed_s=round(time.perf_counter() - t0, 1))
        led.append(rec)
        stats["parsed_now"] += 1
        print(f"  roster {d} {kind}: {len(rows):,} firms, {rec['bytes'] / 1e6:.1f} MB, missing {len(report['missing'])} cols, "
              f"{rec['elapsed_s']} s", flush=True)
    return stats


def _write_roster(rows: list[dict[str, str | None]], d: dt.date, kind: str, av: dt.datetime, basis: str, vr: bool,
                  source_file: str, dest: Path) -> None:
    import pyarrow as pa
    import pyarrow.parquet as pq

    cols: dict[str, Any] = {"roster_date": pa.array([d] * len(rows), pa.date32()),
                            "roster_kind": pa.array([kind] * len(rows), pa.string())}
    for f in ROSTER_FIELDS:
        cols[f] = pa.array([r.get(f) for r in rows], pa.string())
    cols["available_at"] = pa.array([av] * len(rows), pa.timestamp("us", tz="UTC"))
    cols["available_basis"] = pa.array([basis] * len(rows), pa.string())
    cols["vintage_risk"] = pa.array([vr] * len(rows), pa.bool_())
    cols["source_file"] = pa.array([source_file] * len(rows), pa.string())
    tmp = dest.with_name(dest.name + ".partial")
    pq.write_table(pa.table(cols), tmp, compression="zstd")
    os.replace(tmp, dest)


def bulk_directory(part: str) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Central directory of one bulk zip (HEAD + one ranged GET of the tail; cached in the ledger)."""
    led = ledger()
    key = f"bulk_dir:{part}"
    prev = led.latest().get(key)
    if prev and prev.get("status") == "ok":
        return prev["entries"], prev
    entries, info = D.zip_directory(BULK_ZIPS[part])
    rec = {"key": key, "kind": "bulk_directory", "url": BULK_ZIPS[part], "http_status": 206, "fetched_at": S.utc_now(),
           **info, "entries": entries, "status": "ok"}
    led.append(rec)
    return entries, rec


def fetch_bulk(tables: list[str] | None = None) -> dict[str, Any]:
    """Range-fetch, inflate and parse the :data:`BULK_MEMBERS` tables not yet parsed (one member at a time)."""
    led = ledger()
    have = led.latest()
    out_dir = adv_dir() / "bulk"
    out_dir.mkdir(parents=True, exist_ok=True)
    stats: dict[str, Any] = {}
    for table in tables or list(BULK_MEMBERS):
        part, base = BULK_MEMBERS[table]
        key = f"bulk:{table}"
        dest = out_dir / f"{table}.parquet"
        prev = have.get(key)
        if prev and prev.get("status") == "parsed" and dest.exists():
            stats[table] = "skipped"
            continue
        entries, drec = bulk_directory(part)
        entry = next(e for e in entries if e["name"].rsplit("/", 1)[-1].startswith(base + "_"))
        S.require_free(entry["usize"] / 1024 ** 3 * 2 + 1.0)
        lo, hi = D.member_range(entry, drec["zip_bytes"])
        t0 = time.perf_counter()
        raw = RAW_DIR / "bulk" / f"{part}__{entry['name'].rsplit('/', 1)[-1]}.range"
        st, _, hdr = D.get_range(BULK_ZIPS[part], lo, hi, dest=raw)
        rec: dict[str, Any] = {"key": key, "kind": "bulk_member", "url": BULK_ZIPS[part], "member": entry["name"],
                               "range": f"bytes={lo}-{hi}", "http_status": st, "bytes": int(hdr.get("x-bytes", 0)),
                               "sha256": hdr.get("x-sha256"), "fetched_at": S.utc_now(),
                               "zip_last_modified": drec.get("last_modified"), "content_range": hdr.get("content-range")}
        if st != 206:
            rec["status"] = "http_error"
            led.append(rec)
            raise RuntimeError(f"{table}: ranged GET {st}")
        csv_path = S.tmp_dir("adv") / entry["name"].rsplit("/", 1)[-1]
        n_bytes = D.inflate_range_file(raw, entry, csv_path)
        rows = _csv_to_parquet(csv_path, dest)
        csv_path.unlink()
        raw.unlink()
        rec.update(status="parsed", inflated_bytes=n_bytes, crc32=f"{entry['crc']:08x}", rows=rows, parquet=dest.name,
                   parquet_sha256=C.sha256_file(dest), raw_deleted=True, elapsed_s=round(time.perf_counter() - t0, 1))
        led.append(rec)
        stats[table] = rows
        print(f"  bulk {table}: {rows:,} rows, {entry['csize'] / 1e6:.1f} MB deflated, {rec['elapsed_s']} s", flush=True)
    return stats


def _csv_to_parquet(src: Path, dest: Path) -> int:
    """Bulk CSV -> Parquet, all columns VARCHAR with snake-cased names; pyarrow streaming (no DuckDB)."""
    res = D.csv_to_parquet_stream(src, dest, delimiter=",", quoted=True, encoding="latin1", rename=snake)
    if res["skipped_rows"]:
        print(f"    {src.name}: {res['skipped_rows']} unparseable rows skipped", flush=True)
    return int(res["rows"])


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("fetch-rosters")
    r.add_argument("--start", type=dt.date.fromisoformat, default=ROSTER_START)
    r.add_argument("--end", type=dt.date.fromisoformat)
    r.add_argument("--limit", type=int)
    r.add_argument("--months", type=lambda v: tuple(int(x) for x in v.split(",")),
                   help="comma-separated calendar months to keep, e.g. 2,5,8,11")
    b = sub.add_parser("fetch-bulk")
    b.add_argument("--table", action="append", choices=sorted(BULK_MEMBERS))
    sub.add_parser("fetch-all", help="fetch-bulk (every table) then fetch-rosters (default window)")
    args = ap.parse_args(argv)
    if args.cmd == "fetch-rosters":
        D.start_memory_trace()
        print(json.dumps(fetch_rosters(args.start, args.end, args.limit, args.months)), flush=True)
        print(json.dumps({"peak_memory_gb": D.peak_memory_gb(),
                          "run": "unguarded per C-1 (network landing, stdlib csv / openpyxl + pyarrow, no DuckDB)"}), flush=True)
    elif args.cmd == "fetch-bulk":
        D.start_memory_trace()
        print(json.dumps(fetch_bulk(args.table)), flush=True)
        print(json.dumps({"peak_memory_gb": D.peak_memory_gb(),
                          "run": "unguarded per C-1 (network landing, pyarrow streaming CSV, no DuckDB)"}), flush=True)
    else:
        print(json.dumps(fetch_bulk(None)), flush=True)
        print(json.dumps(fetch_rosters()), flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
